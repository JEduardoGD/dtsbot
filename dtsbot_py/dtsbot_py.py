"""Punto de entrada para la generación de archivos DTS."""

import argparse
from datetime import date, datetime
import re
import shutil
import sys
from pathlib import Path

from docx import Document


def iter_paragraphs(contenedor, vistos=None):
    if vistos is None:
        vistos = set()
    for parrafo in contenedor.paragraphs:
        identidad = parrafo._p
        if identidad not in vistos:
            vistos.add(identidad)
            yield parrafo
    for tabla in contenedor.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                yield from iter_paragraphs(celda, vistos)


def reemplazar_marcadores(parrafo, marcadores):
    texto = "".join(run.text for run in parrafo.runs)
    ocurrencias = []
    for marcador, valor in marcadores.items():
        inicio = 0
        while (inicio := texto.find(marcador, inicio)) != -1:
            ocurrencias.append((inicio, inicio + len(marcador), marcador, valor))
            inicio += len(marcador)

    for inicio, fin, marcador, valor in sorted(ocurrencias, reverse=True):
        runs = parrafo.runs
        indice_inicio = indice_fin = None
        offset_inicio = offset_fin = 0
        posicion = 0
        for indice, run in enumerate(runs):
            siguiente = posicion + len(run.text)
            if indice_inicio is None and posicion <= inicio < siguiente:
                indice_inicio = indice
                offset_inicio = inicio - posicion
            if posicion <= fin - 1 < siguiente:
                indice_fin = indice
                offset_fin = fin - posicion
                break
            posicion = siguiente

        run_inicio = runs[indice_inicio]
        run_fin = runs[indice_fin]
        run_inicio.text = run_inicio.text[:offset_inicio] + valor + run_fin.text[offset_fin:]
        for indice in range(indice_inicio + 1, indice_fin + 1):
            runs[indice].text = ""

    return {marcador: sum(1 for ocurrencia in ocurrencias if ocurrencia[2] == marcador)
            for marcador in marcadores}


def main(argv=None):
    arguments = sys.argv[1:] if argv is None else argv
    if not arguments:
        print("Error: debes indicar el archivo TXT de propiedades.", file=sys.stderr)
        return 1

    parser = argparse.ArgumentParser(description="Genera archivos DTS de Word.")
    parser.add_argument("propiedades", help="Archivo TXT con las propiedades")
    args = parser.parse_args(arguments)
    ruta_propiedades = Path(args.propiedades)

    directorio = re.fullmatch(r"INCS(\d+)", ruta_propiedades.parent.name)
    nombre_archivo = re.fullmatch(r"INCS(\d+)", ruta_propiedades.stem)
    if not directorio:
        print("Error: el directorio debe llamarse INCS seguido del número de incidente.", file=sys.stderr)
        return 1
    if not nombre_archivo:
        print("Error: el archivo debe llamarse INCS seguido del número de incidente y tener extensión .txt.", file=sys.stderr)
        return 1
    numero_ruta = directorio.group(1).lstrip("0") or "0"
    numero_archivo = nombre_archivo.group(1).lstrip("0") or "0"
    if numero_ruta != numero_archivo:
        print("Error: el número de incidente del directorio no coincide con el del archivo.", file=sys.stderr)
        return 1

    try:
        with ruta_propiedades.open(encoding="utf-8") as archivo:
            propiedades = []
            for linea in archivo:
                linea = linea.strip()
                if not linea:
                    continue
                separador = "=" if "=" in linea else ":" if ":" in linea else None
                if separador:
                    propiedades.append(linea.split(separador, 1))
    except OSError as error:
        print(f"Error al leer el archivo de propiedades: {error}", file=sys.stderr)
        return 1

    propiedades_requeridas = ("número", "IDENTIFICACION_DEL_PROBLEMA", "FECHA_ANALISIS")
    valores_requeridos = {}
    for requerida in propiedades_requeridas:
        valores = [
            partes[1].strip()
            for partes in propiedades
            if len(partes) == 2 and partes[0].strip() == requerida
        ]
        if not valores:
            print(f"Error: falta la propiedad requerida {requerida}.", file=sys.stderr)
            return 1
        if len(valores) > 1:
            print(f"Error: la propiedad {requerida} aparece más de una vez.", file=sys.stderr)
            return 1
        if not valores[0]:
            print(f"Error: la propiedad {requerida} no puede estar vacía.", file=sys.stderr)
            return 1
        valores_requeridos[requerida] = valores[0]

    numero_propiedad = valores_requeridos["número"]
    numero_propiedad_normalizado = numero_propiedad.lstrip("0") or "0"
    if not re.fullmatch(r"\d+", numero_propiedad) or numero_propiedad_normalizado != numero_ruta:
        print("Error: la propiedad número no coincide con el número de la ruta y el archivo.", file=sys.stderr)
        return 1

    fecha_analisis = valores_requeridos["FECHA_ANALISIS"]
    if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", fecha_analisis):
        print("Error: FECHA_ANALISIS debe tener el formato DD/MM/AAAA.", file=sys.stderr)
        return 1
    try:
        fecha = datetime.strptime(fecha_analisis, "%d/%m/%Y").date()
    except ValueError:
        print("Error: FECHA_ANALISIS no es una fecha válida.", file=sys.stderr)
        return 1
    if fecha != date.today():
        print(
            f"Advertencia: FECHA_ANALISIS ({fecha_analisis}) no coincide con la fecha actual ({date.today():%d/%m/%Y})."
        )

    ejemplo = Path(__file__).resolve().parent / "examples" / "NNNNNNN_DTS.docx"
    destino = ruta_propiedades.parent / f"{numero_ruta}_DTS.docx"
    try:
        shutil.copy2(ejemplo, destino)
        documento = Document(destino)
        marcadores = {
            "[<número>]": valores_requeridos["número"],
            "[IDENTIFICACION_DEL_PROBLEMA]": valores_requeridos["IDENTIFICACION_DEL_PROBLEMA"],
            "[FECHA_ANALISIS]": fecha_analisis,
        }
        encontrados = {marcador: 0 for marcador in marcadores}
        secciones = []
        for seccion in documento.sections:
            secciones.extend((
                seccion.header,
                seccion.first_page_header,
                seccion.even_page_header,
                seccion.footer,
                seccion.first_page_footer,
                seccion.even_page_footer,
            ))
        contenedores = [documento, *secciones]
        vistos = set()
        for contenedor in contenedores:
            for parrafo in iter_paragraphs(contenedor, vistos):
                conteos = reemplazar_marcadores(parrafo, marcadores)
                for marcador, cantidad in conteos.items():
                    encontrados[marcador] += cantidad
        sin_marcador = [marcador for marcador, cantidad in encontrados.items() if cantidad == 0]
        if sin_marcador:
            print(f"Error: no se encontraron los marcadores {', '.join(sin_marcador)} en el documento.", file=sys.stderr)
            return 1
        documento.save(destino)
    except Exception as error:
        print(f"Error al procesar el documento DTS: {error}", file=sys.stderr)
        return 1

    print("Propiedades:")
    for partes in propiedades:
        if len(partes) == 2:
            nombre, valor = partes
            print(f"{nombre.strip()}: {valor.strip()}")
    print("Marcadores reemplazados:")
    for marcador, cantidad in encontrados.items():
        print(f"{marcador}: {cantidad}")
    print(f"Documento DTS copiado a: {destino}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
